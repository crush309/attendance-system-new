from io import BytesIO
import json
import os
from pathlib import Path
import socket
import subprocess
import sys
import tempfile
import time
import unittest
from urllib.error import URLError, HTTPError
from urllib.request import Request, urlopen

from openpyxl import Workbook, load_workbook

ROOT = Path(__file__).resolve().parents[1]


class LocalHttpTests(unittest.TestCase):
    SERVERLESS = False

    @classmethod
    def setUpClass(cls):
        cls.temp = tempfile.TemporaryDirectory()
        env = dict(os.environ)
        env.pop('DATABASE_URL', None)
        env['SQLITE_DB_PATH'] = str(Path(cls.temp.name) / 'test.db')
        if cls.SERVERLESS:
            env.update(VERCEL='1', TEMP=cls.temp.name, TMP=cls.temp.name, TMPDIR=cls.temp.name)
        with socket.socket() as sock:
            sock.bind(('127.0.0.1', 0))
            port = sock.getsockname()[1]
        cls.base = f'http://127.0.0.1:{port}'
        cls.log = open(Path(cls.temp.name) / 'server.log', 'w+b')
        entrypoint = 'app:app' if cls.SERVERLESS else 'backend.app:app'
        cls.process = subprocess.Popen([sys.executable, '-m', 'uvicorn', entrypoint,
                                       '--host', '127.0.0.1', '--port', str(port)],
                                      cwd=ROOT, env=env, stdout=cls.log, stderr=cls.log)
        for _ in range(100):
            try:
                with urlopen(cls.base, timeout=1) as response:
                    if response.status == 200:
                        return
            except URLError:
                if cls.process.poll() is not None:
                    break
                time.sleep(0.1)
        cls.process.terminate()
        cls.process.wait(timeout=10)
        cls.log.seek(0)
        message = cls.log.read().decode('utf-8', errors='replace')
        cls.log.close()
        cls.temp.cleanup()
        raise AssertionError('Local server did not start: ' + message)

    @classmethod
    def tearDownClass(cls):
        cls.process.terminate()
        cls.process.wait(timeout=10)
        cls.log.close()
        cls.temp.cleanup()

    def request(self, path, data=None, token=None, content_type='application/json'):
        headers = {'Content-Type': content_type}
        if token:
            headers['Authorization'] = 'Bearer ' + token
        if isinstance(data, dict):
            data = json.dumps(data).encode('utf-8')
        with urlopen(Request(self.base + path, data=data, headers=headers), timeout=10) as response:
            return response.read()

    def test_home_login_upload_export(self):
        self.assertIn('考勤数据分析系统', self.request('/').decode('utf-8'))
        with self.assertRaises(HTTPError) as denied:
            self.request('/api/login', {'username': 'admin', 'password': 'wrong'})
        self.assertEqual(denied.exception.code, 401)
        token = json.loads(self.request('/api/login', {'username': 'admin', 'password': 'admin123'}))['token']
        self.assertEqual(json.loads(self.request('/api/me', token=token))['role'], 'admin')
        wb = Workbook()
        ws = wb.active
        ws.append(['统计日期:2026/10/01-2026/10/02'])
        ws.append(['工号', '姓名', '部门', '1', '2'])
        ws.append(['001', '测试员工', '测试部门', '08:30\n14:00', None])
        buf = BytesIO()
        wb.save(buf)
        boundary = 'attendance-test-boundary'
        body = (f'--{boundary}\r\nContent-Disposition: form-data; name="file"; filename="test.xlsx"\r\n'
                'Content-Type: application/vnd.openxmlformats-officedocument.spreadsheetml.sheet\r\n\r\n').encode()
        body += buf.getvalue() + f'\r\n--{boundary}--\r\n'.encode()
        result = json.loads(self.request('/api/upload', body, token, 'multipart/form-data; boundary=' + boundary))
        self.assertEqual(result['total_days'], 2)
        record = result['records'][0]
        self.assertEqual(record['attendance_days'], 1)
        self.assertEqual(record['absent_days'], 1)
        self.assertEqual(record['total_punches'], 2)
        exported = self.request('/api/export', {'records': result['records'], 'include_detail': True}, token)
        output = load_workbook(BytesIO(exported), data_only=True)
        self.assertEqual(output.sheetnames, ['考勤统计', '考勤明细', '考勤透视'])
        self.assertEqual(output['考勤统计'].cell(2, 2).value, '测试员工')
        self.assertEqual(output['考勤统计'].cell(2, 5).value, 1)
        self.assertEqual(output['考勤统计'].cell(2, 6).value, 1)
        self.assertEqual(output['考勤统计'].cell(2, 8).value, 2)
        self.assertEqual(output['考勤明细'].cell(2, 4).value, '10月1日')
        self.assertEqual(output['考勤透视'].cell(3, 4).value, 1)

    def upload_rows(self, rows, token):
        wb = Workbook()
        ws = wb.active
        ws.append(['统计日期:2026/09/01-2026/09/02'])
        ws.append(['工号', '姓名', '部门', '1', '2'])
        for eid, name in rows:
            ws.append([eid, name, '测试部门', '08:30\n14:00', None])
        buf = BytesIO()
        wb.save(buf)
        boundary = 'group-name-test'
        body = (f'--{boundary}\r\nContent-Disposition: form-data; name="file"; filename="groups.xlsx"\r\n'
                'Content-Type: application/octet-stream\r\n\r\n').encode()
        body += buf.getvalue() + f'\r\n--{boundary}--\r\n'.encode()
        return json.loads(self.request('/api/upload', body, token, 'multipart/form-data; boundary=' + boundary))

    def test_groups_follow_names_when_ids_change_in_upload_and_export(self):
        token = json.loads(self.request('/api/login', {'username': 'admin', 'password': 'admin123'}))['token']
        group = json.loads(self.request('/api/groups', {'name': '按姓名验证组'}, token))
        gid = group['id']
        self.upload_rows([('001', '张三'), ('002', '李四')], token)
        assigned = json.loads(self.request('/api/groups/assign',
                                          {'employee_names': [' 张三 ', '张三'], 'group_id': gid}, token))
        self.assertEqual(assigned['count'], 1)
        result = self.upload_rows([('099', '张三'), ('001', '王五'), ('100', '李四')], token)
        self.assertEqual(result['records'][0]['groups'], [{'id': gid, 'name': '按姓名验证组'}])
        self.assertEqual(result['records'][1]['groups'], [])
        self.assertEqual(json.loads(self.request(f'/api/groups/{gid}/members', token=token)), ['张三'])
        mapping = json.loads(self.request('/api/groups/mapping', token=token))
        self.assertEqual(mapping['张三'], [{'id': gid, 'name': '按姓名验证组'}])
        exported = self.request('/api/export', {'records': result['records'], 'group_id': gid}, token)
        output = load_workbook(BytesIO(exported), data_only=True)
        summary = output.worksheets[0]
        self.assertEqual(summary.max_row, 2)
        self.assertEqual(summary.cell(2, 1).value, '099')
        self.assertEqual(summary.cell(2, 2).value, '张三')
        self.assertEqual(summary.cell(2, 4).value, '按姓名验证组')
        self.request('/api/groups/remove', {'employee_name': ' 张三 ', 'group_id': gid}, token)
        self.assertEqual(json.loads(self.request(f'/api/groups/{gid}/members', token=token)), [])
        self.assertEqual(self.upload_rows([('001', '张三')], token)['records'][0]['groups'], [])

    def test_old_id_payloads_and_blank_names_are_rejected(self):
        token = json.loads(self.request('/api/login', {'username': 'admin', 'password': 'admin123'}))['token']
        gid = json.loads(self.request('/api/groups', {'name': '校验组'}, token))['id']
        for payload in ({'emp_ids': ['001'], 'group_id': gid},
                        {'employee_names': ['   '], 'group_id': gid},
                        {'employee_names': [], 'group_id': gid}):
            with self.assertRaises(HTTPError) as invalid:
                self.request('/api/groups/assign', payload, token)
            self.assertEqual(invalid.exception.code, 422)


class ServerlessHttpTests(LocalHttpTests):
    SERVERLESS = True


if __name__ == '__main__':
    unittest.main()
